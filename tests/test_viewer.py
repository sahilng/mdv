import io
import subprocess
from unittest.mock import AsyncMock, Mock

import pytest
from rich.cells import cell_len
from textual import events
from textual._xterm_parser import XTermParser
from textual.command import CommandPalette
from textual.selection import SELECT_ALL
from textual.widgets import Input, Markdown, MarkdownViewer, Static, TextArea
from textual.widgets._markdown import MarkdownTableOfContents
from textual.widgets._footer import FooterKey
from textual.widgets import Tree

from mdv.app import Viewer
from mdv.editor import MarkdownEditor
from mdv.preview import AlignedPreview
from mdv.rendered import SOURCE
from mdv.cli import main
from mdv.document import load_document


async def settle_preview(app, pilot):
    await pilot.pause(0.15)
    # Rendering is queued after layout and no longer blocks Changed messages.
    for _ in range(10):
        await pilot.pause()
        await app.workers.wait_for_complete()
        await pilot.pause()
        if not app._refreshing_preview and not app._preview_running and not app._preview_scheduled:
            return
    raise AssertionError("Preview did not settle")


@pytest.mark.parametrize("key", ["super+c", "ctrl+shift+c"])
@pytest.mark.parametrize("mode", ["read", "live", "split"])
async def test_copy_selected_text(tmp_path, monkeypatch, key, mode):
    path = tmp_path / "copy.md"
    path.write_text("Selected café\n\nOther text\n")
    app = Viewer(path, start_editing=mode != "read",
                 live_edit=None if mode == "read" else mode == "live")
    native_copy = Mock()
    monkeypatch.setattr("mdv.app.subprocess.run", native_copy)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        if mode == "read":
            paragraph = app.query_one(MarkdownViewer).document.query("MarkdownParagraph").first()
            app.screen.selections = {paragraph: SELECT_ALL}
        else:
            editor = app.query_one(TextArea)
            editor.move_cursor((0, 0))
            editor.move_cursor((0, len("Selected café")), select=True)
        exit_app = Mock()
        monkeypatch.setattr(app, "exit", exit_app)
        await pilot.press(key)
        assert app.clipboard == "Selected café"
        assert not app.dirty
        exit_app.assert_not_called()
        if native_copy.called:
            assert native_copy.call_args.kwargs["input"] == "Selected café".encode("utf-8")


async def test_copy_without_selection_preserves_clipboard(tmp_path, monkeypatch):
    path = tmp_path / "copy.md"
    path.write_text("Text\n")
    app = Viewer(path, start_editing=True)
    copy = Mock()
    monkeypatch.setattr(app, "copy_to_clipboard", copy)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.press("super+c", "ctrl+shift+c")
        copy.assert_not_called()


async def test_copy_in_command_palette(tmp_path, monkeypatch):
    path = tmp_path / "copy.md"
    path.write_text("Text\n")
    app = Viewer(path)
    monkeypatch.setattr("mdv.app.subprocess.run", Mock())
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.press("ctrl+p")
        search = app.screen.query_one(Input)
        search.value = "theme"
        search.select_all()
        await pilot.press("super+c")
        assert app.clipboard == "theme"
        assert search.value == "theme"
        assert CommandPalette.is_open(app)


@pytest.mark.parametrize("mode", ["read", "live", "split", "preview"])
@pytest.mark.parametrize("platform", ["darwin", "linux"])
@pytest.mark.parametrize("trigger", ["key", "footer"])
async def test_mouse_selection_copy(tmp_path, monkeypatch, mode, platform, trigger):
    path = tmp_path / "copy.md"
    path.write_text("Selected café\n\nOther text\n")
    app = Viewer(path, start_editing=mode != "read",
                 live_edit=None if mode == "read" else mode == "live")
    native_copy = Mock()
    monkeypatch.setattr("mdv.app.subprocess.run", native_copy)
    monkeypatch.setattr("mdv.app.sys.platform", platform)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        if mode != "read":
            await settle_preview(app, pilot)
        else:
            await pilot.pause()
        if mode == "read":
            target = app.query_one(MarkdownViewer).document.query("MarkdownParagraph").first()
            start = (0, 0)
        elif mode == "preview":
            target = app.query_one(AlignedPreview)
            start = (target.gutter.left, target.gutter.top)
        else:
            target = app.query_one(TextArea)
            start = (target.gutter.left + target.gutter_width, target.gutter.top)
        await pilot.mouse_down(target, offset=start)
        await pilot.hover(target, offset=(start[0] + 8, start[1]))
        assert app.clipboard == ""
        await pilot.mouse_up(target, offset=(start[0] + 8, start[1]))
        await pilot.pause()
        assert app.clipboard == ""
        native_copy.assert_not_called()
        if trigger == "key":
            await pilot.press("ctrl+shift+c")
        else:
            copy = next(key for key in app.query(FooterKey) if key.description == "Copy")
            assert await pilot.click(copy)
        expected = "Selected " if mode in {"read", "preview"} else "Selected"
        assert app.clipboard == expected
        if platform == "darwin":
            assert native_copy.call_args.kwargs["input"] == expected.encode()
        else:
            native_copy.assert_not_called()
        assert not app.dirty
        native_copy.reset_mock()
        await pilot.click(target, offset=start)
        native_copy.assert_not_called()


async def test_macos_keyboard_selection_does_not_copy_automatically(tmp_path, monkeypatch):
    path = tmp_path / "copy.md"
    path.write_text("Selected café\n")
    app = Viewer(path, start_editing=True)
    native_copy = Mock()
    monkeypatch.setattr("mdv.app.subprocess.run", native_copy)
    monkeypatch.setattr("mdv.app.sys.platform", "darwin")
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.press("shift+right", "shift+right")
        assert app.clipboard == ""
        native_copy.assert_not_called()
        await pilot.press("ctrl+shift+c")
        assert app.clipboard == "Se"
        assert native_copy.call_args.kwargs["input"] == b"Se"


@pytest.mark.parametrize("source,code", [
    ("```python\nprint('café')  \n\n```\n", "print('café')  \n\n"),
    ("    print('indented')\n", "print('indented')\n"),
    ("```\n```\n", ""),
])
async def test_code_block_copy_button(tmp_path, monkeypatch, source, code):
    path = tmp_path / "code.md"
    path.write_text(source)
    app = Viewer(path)
    native_copy = Mock()
    monkeypatch.setattr("mdv.app.subprocess.run", native_copy)
    monkeypatch.setattr("mdv.app.sys.platform", "darwin")
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        button = app.query_one(".copy-code")
        block = button.parent
        assert button.region.right == block.content_region.right
        assert button.region.y == block.content_region.y
        native_copy.assert_not_called()
        assert await pilot.click(button)
        assert app.clipboard == code
        assert native_copy.call_args.kwargs["input"] == code.encode("utf-8")
        assert path.read_text() == source


async def test_code_buttons_copy_their_own_block_after_reload(tmp_path, monkeypatch):
    path = tmp_path / "code.md"
    path.write_text("```\nfirst\n```\n\n```\nsecond\n```\n")
    app = Viewer(path)
    monkeypatch.setattr("mdv.app.subprocess.run", Mock())
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        buttons = list(app.query(".copy-code"))
        assert len(buttons) == 2
        await pilot.click(buttons[1])
        assert app.clipboard == "second\n"
        path.write_text("```\nupdated\n```\n")
        await pilot.press("r")
        await app.workers.wait_for_complete()
        await pilot.pause()
        assert len(app.query(".copy-code")) == 1
        await pilot.click(".copy-code")
        assert app.clipboard == "updated\n"


async def test_code_copy_confirmation_resets_after_latest_click(tmp_path, monkeypatch):
    path = tmp_path / "code.md"
    path.write_text("```\nfirst\n```\n\n```\nsecond\n```\n")
    app = Viewer(path)
    monkeypatch.setattr("mdv.app.subprocess.run", Mock())
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        buttons = list(app.query(".copy-code"))
        await pilot.pause()
        await pilot.click(buttons[0])
        assert str(buttons[0].label) == "✓ Copy"
        assert str(buttons[1].label) == "⧉ Copy"
        await pilot.pause(0.8)
        await pilot.click(buttons[0])
        await pilot.pause(0.8)
        assert str(buttons[0].label) == "✓ Copy"
        await pilot.pause(0.8)
        assert str(buttons[0].label) == "⧉ Copy"


async def test_code_copy_failure_keeps_original_icon(tmp_path, monkeypatch):
    path = tmp_path / "code.md"
    path.write_text("```\ncode\n```\n")
    app = Viewer(path)
    monkeypatch.setattr("mdv.app.sys.platform", "darwin")
    monkeypatch.setattr("mdv.app.subprocess.run", Mock(side_effect=OSError("unavailable")))
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        button = app.query_one(".copy-code")
        await pilot.pause()
        await pilot.click(button)
        assert str(button.label) == "⧉ Copy"


@pytest.mark.parametrize("theme", ["textual-dark", "textual-light", "atom-one-dark"])
async def test_code_copy_icon_stays_inside_block_on_hover_and_resize(tmp_path, monkeypatch, theme):
    path = tmp_path / "code.md"
    code = "echo first\n" + "long_code_" * 20 + "\n"
    path.write_text(f"```sh\n{code}```\n")
    monkeypatch.setenv("TEXTUAL_THEME", theme)
    app = Viewer(path)
    async with app.run_test(size=(100, 20)) as pilot:
        await app.workers.wait_for_complete()
        for width in (100, 50, 120, 30, 80):
            await pilot.resize_terminal(width, 20)
            await pilot.pause()
            button = app.query_one(".copy-code")
            block = button.parent
            await pilot.hover(button)
            await pilot.pause()
            assert button.region.right == block.content_region.right
            assert button.region.y == block.content_region.y
            assert button.content_region.height == 1
            assert block.region.contains_region(button.region)
            rows = app.screen._compositor.render_strips()
            # Button's default line padding can paint past its allocated width
            # even when its reported region is correctly inside the block.
            assert all(cell_len(row.text) == width for row in rows)
            assert "echo first" in rows[button.region.y + 2].text
            assert str(button.label) in rows[button.region.y].text
            block.scroll_to(x=15, animate=False, immediate=True)
            await pilot.pause()
            assert button.region.right == block.content_region.right
            assert str(button.label) in app.screen._compositor.render_strips()[button.region.y].text
            assert await pilot.click(button)
            assert app.clipboard == code
            assert all(cell_len(row.text) == width for row in app.screen._compositor.render_strips())
            block.scroll_to(x=0, animate=False, immediate=True)


async def test_narrow_code_block_scrollbar_and_keyboard_navigation(tmp_path):
    path = tmp_path / "wide-code.md"
    code = "START_" + "middle_" * 20 + "END_TOKEN\n"
    path.write_text(f"```sh\n{code}```\n\n```\nshort\n```\n\n" + "Normal prose wraps. " * 8)
    app = Viewer(path)
    async with app.run_test(size=(40, 24)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        long, short = app.query("CopyableMarkdownFence")
        assert long.show_horizontal_scrollbar
        assert long.horizontal_scrollbar.region.height == 1
        assert not short.show_horizontal_scrollbar
        # Keep source lines intact; the scrollbar adds one row to the panel.
        assert long.query_one("#code-content").content_region.height == 1
        prose = app.query_one(MarkdownViewer).document.query("MarkdownParagraph").first()
        assert prose.content_region.height > 1
        assert await pilot.click(long.query_one("#code-content"), offset=(2, 1))
        assert long.has_focus
        await pilot.press("right")
        await pilot.pause(0.2)
        assert long.scroll_x > 0
        after_key = long.scroll_x
        bar = long.horizontal_scrollbar
        assert await pilot.mouse_down(bar, offset=(1, 0))
        await pilot.hover(bar, offset=(bar.region.width - 1, 0))
        await pilot.mouse_up(bar, offset=(bar.region.width - 1, 0))
        await pilot.pause(0.2)
        assert long.scroll_x > after_key
        # The thumb uses solid cells, with no fractional glyphs at its ends.
        bar_row = app.screen._compositor.render_strips()[bar.region.y]
        assert not bar_row.crop(bar.region.x, bar.region.right).text.strip()
        assert "END_TOKEN" in app.screen._compositor.render_strips()[long.region.y + 2].text
        button = long.query_one(".copy-code")
        assert button.region.right == long.content_region.right
        assert await pilot.click(button)
        assert app.clipboard == code
        await pilot.resize_terminal(220, 24)
        await pilot.pause()
        assert not long.show_horizontal_scrollbar
        assert long.scroll_x == 0


@pytest.mark.parametrize("shifted", [False, True])
async def test_trackpad_horizontal_scroll_over_code(tmp_path, shifted):
    path = tmp_path / "scroll.md"
    path.write_text("```\n" + "wide_line_" * 25 + "\n```\n\n" + "Paragraph\n\n" * 25)
    app = Viewer(path)
    async with app.run_test(size=(40, 18)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        viewer = app.query_one(MarkdownViewer)
        block = app.query_one("CopyableMarkdownFence")
        label = block.query_one("#code-content")
        x, y = label.content_region.offset
        # SGR 66/67 are horizontal wheel reports; 68/69 are Shift+vertical.
        parser = XTermParser()
        event = parser.parse_mouse_code(f"\x1b[<{69 if shifted else 67};{x + 1};{y + 1}M")
        assert isinstance(event, events.MouseScrollDown if shifted else events.MouseScrollRight)
        app.post_message(event)
        await pilot.pause(0.2)
        assert block.scroll_x > 0
        assert viewer.scroll_y == 0
        event = parser.parse_mouse_code(f"\x1b[<{68 if shifted else 66};{x + 1};{y + 1}M")
        app.post_message(event)
        await pilot.pause(0.2)
        assert block.scroll_x == 0
        assert viewer.scroll_y == 0


@pytest.mark.parametrize("editing", [False, True])
async def test_vertical_scrollbar_edges_after_scroll_and_resize(tmp_path, editing):
    path = tmp_path / "vertical.md"
    path.write_text("\n\n".join(f"Paragraph {index}" for index in range(80)))
    app = Viewer(path, start_editing=editing)
    async with app.run_test(size=(50, 20)) as pilot:
        await app.workers.wait_for_complete()
        if editing:
            await settle_preview(app, pilot)
        else:
            await pilot.pause()
        target = app.query_one(MarkdownEditor if editing else MarkdownViewer)
        for width, height, position in [
            (50, 20, 0), (50, 20, 7), (20, 16, 13), (12, 12, 19),
            (8, 10, 30), (70, 24, 13),
        ]:
            await pilot.resize_terminal(width, height)
            target.scroll_to(y=position, animate=False, immediate=True)
            await pilot.pause()
            bar = target.vertical_scrollbar
            assert target.show_vertical_scrollbar
            await pilot.hover(bar, offset=(0, 0))
            rows = app.screen._compositor.render_strips()
            for row in rows[bar.region.y:bar.region.bottom]:
                assert not row.crop(bar.region.x, bar.region.right).text.strip()
                assert cell_len(row.text) == width


@pytest.mark.parametrize("at_edge", [False, True])
async def test_shift_scroll_momentum_keeps_horizontal_axis(tmp_path, at_edge):
    path = tmp_path / "momentum.md"
    path.write_text("```\n" + "wide_line_" * 25 + "\n```\n\n" + "Paragraph\n\n" * 30)
    app = Viewer(path)
    async with app.run_test(size=(40, 18)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        viewer = app.query_one(MarkdownViewer)
        block = app.query_one("CopyableMarkdownFence")
        if at_edge:
            block.scroll_to(x=block.max_scroll_x, animate=False, immediate=True)
        x, y = block.query_one("#code-content").content_region.offset
        parser = XTermParser()
        # Trackpad momentum may lose its Shift modifier when Shift is released.
        app.post_message(parser.parse_mouse_code(f"\x1b[<69;{x + 1};{y + 1}M"))
        await pilot.pause(0.03)
        for _ in range(4):
            app.post_message(parser.parse_mouse_code(f"\x1b[<65;{x + 1};{y + 1}M"))
            await pilot.pause(0.03)
        assert viewer.scroll_y == 0
        assert block.scroll_x > 0
        # After the gesture ends, normal vertical scrolling must work again.
        await pilot.pause(0.35)
        app.post_message(parser.parse_mouse_code(f"\x1b[<65;{x + 1};{y + 1}M"))
        await pilot.pause(0.2)
        assert viewer.scroll_y > 0


@pytest.mark.parametrize("wheel_code", [69, 85, 67])
async def test_unmodified_trackpad_scroll_cannot_extend_horizontal_lock(tmp_path, wheel_code):
    path = tmp_path / "scroll.md"
    path.write_text("```\n" + "wide_line_" * 25 + "\n```\n\n" + "Paragraph\n\n" * 30)
    app = Viewer(path)
    async with app.run_test(size=(40, 18)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        viewer = app.query_one(MarkdownViewer)
        block = app.query_one("CopyableMarkdownFence")
        x, y = block.query_one("#code-content").content_region.offset
        parser = XTermParser()
        app.post_message(parser.parse_mouse_code(f"\x1b[<{wheel_code};{x + 1};{y + 1}M"))
        await pilot.pause(0.03)
        assert block.scroll_x > 0
        # No idle gap: ordinary wheel events must still let the lock expire.
        for _ in range(12):
            app.post_message(parser.parse_mouse_code(f"\x1b[<65;{x + 1};{y + 1}M"))
            await pilot.pause(0.03)
        assert not block._horizontal_scroll_active
        assert viewer.scroll_y > 0


@pytest.mark.parametrize("error", [None, OSError("unavailable"), subprocess.TimeoutExpired("pbcopy", 2)])
def test_native_macos_clipboard(tmp_path, monkeypatch, error):
    app = Viewer(tmp_path / "copy.md")
    native_copy = Mock(side_effect=error)
    notify = Mock()
    monkeypatch.setattr("mdv.app.sys.platform", "darwin")
    monkeypatch.setattr("mdv.app.subprocess.run", native_copy)
    monkeypatch.setattr(app, "notify", notify)
    assert app.copy_to_clipboard("café\nsecond line") is (error is None)
    assert app.clipboard == "café\nsecond line"
    native_copy.assert_called_once_with(
        ["/usr/bin/pbcopy"], input="café\nsecond line".encode("utf-8"),
        check=True, timeout=2, capture_output=True,
    )
    assert notify.called is (error is not None)


@pytest.mark.parametrize("existing", [False, True])
async def test_start_in_editor_and_save(tmp_path, existing):
    path = tmp_path / "notes.md"
    if existing:
        path.write_text("# Existing\n")
    app = Viewer(path, start_editing=True)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        editor = app.query_one(TextArea)
        assert app.editing
        assert editor.live_render
        assert editor.has_focus
        assert editor.text == ("# Existing\n" if existing else "")
        assert path.exists() is existing
        editor.load_text("# Saved\n")
        await pilot.press("ctrl+s", "escape")
        assert path.read_text() == "# Saved\n"
        assert not app.editing


async def test_contents_toggle_and_navigation_in_live_editor(tmp_path):
    path = tmp_path / "contents.md"
    path.write_text("# First\n\nText\n\n## Second\n\nMore text\n")
    app = Viewer(path)
    async with app.run_test(size=(80, 20)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        viewer = app.query_one(MarkdownViewer)
        edit_toc = app.query_one("#edit-toc", MarkdownTableOfContents)
        assert not viewer.show_table_of_contents
        await pilot.press("ctrl+t")
        assert viewer.show_table_of_contents
        await pilot.press("e")
        await settle_preview(app, pilot)
        editor = app.query_one(MarkdownEditor)
        assert editor.live_render
        assert edit_toc.display
        await pilot.press("ctrl+t")
        await settle_preview(app, pilot)
        assert not edit_toc.display
        assert not viewer.show_table_of_contents
        await pilot.press("ctrl+t")
        await settle_preview(app, pilot)
        assert edit_toc.display
        tree = edit_toc.query_one(Tree)
        assert len(tree.root.children) == 1
        tree.select_node(tree.root.children[0].children[0])
        await pilot.pause()
        assert editor.cursor_location == (4, 0)
        assert editor.has_focus
        editor.replace("Updated", (4, 3), (4, 9))
        await settle_preview(app, pilot)
        assert edit_toc.table_of_contents[1][1] == "Updated"
        await pilot.press("ctrl+s", "escape")
        assert not app.editing
        assert viewer.show_table_of_contents
        assert viewer.document.table_of_contents[1][1] == "Updated"


async def test_read_edit_transition_keeps_nearby_content(tmp_path):
    path = tmp_path / "position.md"
    path.write_text("\n\n".join(f"## Section {index}\n\nParagraph {index}" for index in range(40)))
    app = Viewer(path)
    async with app.run_test(size=(80, 20)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        viewer = app.query_one(MarkdownViewer)
        heading = next(block for block in viewer.document.children
                       if getattr(block, "source_range", (None,))[0] == 80)
        viewer.scroll_to_widget(heading, top=True, animate=False)
        await pilot.pause()
        await pilot.press("e")
        await settle_preview(app, pilot)
        editor = app.query_one(MarkdownEditor)
        assert editor.cursor_location[0] >= 75
        assert editor.scroll_y > 0
        await pilot.press("escape")
        await pilot.pause()
        assert viewer.scroll_y > 0
        assert app._read_top_source_row(viewer) >= 75


@pytest.mark.parametrize("sidebar", [False, True])
async def test_live_editor_matches_read_margins(tmp_path, sidebar):
    path = tmp_path / "margins.md"
    path.write_text("# Heading\n\nBody\n")
    app = Viewer(path, show_toc=sidebar)
    async with app.run_test(size=(90, 22)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        reader = app.query_one(MarkdownViewer).document
        read_margin = (reader.content_region.x - reader.region.x,
                       reader.content_region.y - reader.region.y)
        await pilot.press("e")
        await settle_preview(app, pilot)
        editor = app.query_one(MarkdownEditor)
        assert editor.live_render
        assert (editor.content_region.x - editor.region.x,
                editor.content_region.y - editor.region.y) == read_margin == (3, 1)
        assert app.query_one("#edit-toc", MarkdownTableOfContents).display is sidebar


async def test_sidebar_shortcut_hidden_in_split_editor(tmp_path):
    path = tmp_path / "split.md"
    path.write_text("# Heading\n")
    app = Viewer(path, start_editing=True)
    async with app.run_test() as pilot:
        await settle_preview(app, pilot)
        assert "ctrl+t" in app.screen.active_bindings
        await pilot.press("ctrl+l")
        await settle_preview(app, pilot)
        assert not app.live_edit
        assert "ctrl+t" not in app.screen.active_bindings
        await pilot.press("ctrl+t")
        assert not app.show_toc
        await pilot.press("ctrl+l")
        await settle_preview(app, pilot)
        assert "ctrl+t" in app.screen.active_bindings
        await pilot.press("ctrl+t")
        assert app.show_toc


async def test_empty_read_view_can_be_edited_and_rendered(tmp_path):
    path = tmp_path / "empty.md"
    path.write_text("")
    app = Viewer(path)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        reader = app.query_one(MarkdownViewer).document
        assert reader.source == ""
        await pilot.press("e")
        assert app.editing
        editor = app.query_one(TextArea)
        editor.load_text("# Heading\n")
        await pilot.press("ctrl+s", "escape")
        assert not app.editing
        assert reader.source == "# Heading\n"
        assert reader.table_of_contents[0][1] == "Heading"


def test_cli_interactive_validation(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("sys.stdin.isatty", lambda: True)
    monkeypatch.setattr("sys.stdout.isatty", lambda: True)
    viewer = Mock()
    monkeypatch.setattr("mdv.app.Viewer", viewer)
    path = tmp_path / "new.md"
    assert main([str(path)]) == 1
    assert "mde" in capsys.readouterr().err
    viewer.assert_not_called()
    assert main(["--edit", str(path)]) == 0
    viewer.assert_called_once_with(path, show_toc=False, start_editing=True)
    viewer.return_value.run.assert_called_once()
    assert not path.exists()
    viewer.reset_mock()
    assert main(["--toc", "--edit", str(path)]) == 0
    viewer.assert_called_once_with(path, show_toc=True, start_editing=True)
    viewer.reset_mock()
    for invalid in (tmp_path, tmp_path / "new.pdf", tmp_path / "missing" / "new.md"):
        assert main(["--edit", str(invalid)]) == 1
    viewer.assert_not_called()


def test_edit_requires_path_and_terminal(tmp_path, monkeypatch, capsys):
    for args in (["--edit"], ["--edit", "-"], ["--edit", "--print", "file.md"]):
        with pytest.raises(SystemExit) as error:
            main(args)
        assert error.value.code == 2
    capsys.readouterr()
    monkeypatch.setattr("sys.stdin", io.StringIO(""))
    assert main(["--edit", str(tmp_path / "new.md")]) == 1
    assert "interactive terminal" in capsys.readouterr().err


@pytest.mark.parametrize("entry, flag", [("print_main", "--print"), ("edit_main", "--edit")])
def test_shortcuts(entry, flag, monkeypatch):
    from mdv import cli

    main_mock = Mock(return_value=0)
    monkeypatch.setattr(cli, "main", main_mock)
    monkeypatch.setattr("sys.argv", ["shortcut", "notes.md", "--no-toc"])
    assert getattr(cli, entry)() == 0
    main_mock.assert_called_once_with([flag, "notes.md", "--no-toc"])


@pytest.mark.parametrize("dirty", [False, True])
async def test_escape_palette_preserves_edit_mode(tmp_path, dirty):
    path = tmp_path / "edit.md"
    path.write_text("# Original\n")
    app = Viewer(path, show_toc=True)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.press("e")
        editor = app.query_one(TextArea)
        viewer = app.query_one(MarkdownViewer)
        document_screen = app.screen
        if dirty:
            editor.load_text("# Changed\n")
        await pilot.press("ctrl+p")
        assert isinstance(app.screen, CommandPalette)
        await pilot.press("escape")
        assert app.screen is document_screen
        assert app.editing
        assert document_screen.has_class("editing")
        assert editor.has_focus
        assert viewer.show_table_of_contents
        assert editor.text == ("# Changed\n" if dirty else "# Original\n")
        assert app.dirty is dirty
        await pilot.press("ctrl+d" if dirty else "escape")
        assert not app.editing
        assert not document_screen.has_class("editing")
        assert viewer.show_table_of_contents


@pytest.mark.parametrize("dirty", [False, True])
async def test_quit_directly_from_editor(tmp_path, monkeypatch, dirty):
    path = tmp_path / "edit.md"
    path.write_text("Original")
    app = Viewer(path, start_editing=True)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        exit_app = Mock()
        monkeypatch.setattr(app, "exit", exit_app)
        editor = app.query_one(TextArea)
        if dirty:
            editor.load_text("Changed")
        await pilot.press("ctrl+p", "ctrl+c")
        exit_app.assert_not_called()
        await pilot.press("escape", "ctrl+c")
        if dirty:
            exit_app.assert_not_called()
            assert editor.text == "Changed"
            assert path.read_text() == "Original"
            await pilot.press("ctrl+s", "ctrl+c")
            assert path.read_text() == "Changed"
        exit_app.assert_called_once_with()
        assert app.editing


async def test_ctrl_c_replaces_q_in_read_mode(tmp_path, monkeypatch):
    path = tmp_path / "read.md"
    path.write_text("# Heading\n")
    app = Viewer(path)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        exit_app = Mock()
        monkeypatch.setattr(app, "exit", exit_app)
        await pilot.press("q")
        exit_app.assert_not_called()
        await pilot.press("ctrl+c")
        exit_app.assert_called_once_with()


async def test_edit_preview_save_and_discard(tmp_path):
    path = tmp_path / "edit.md"
    path.write_text("# Original\n")
    app = Viewer(path, show_toc=True, live_edit=False)
    async with app.run_test(size=(100, 30)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.press("e")
        editor = app.query_one(TextArea)
        viewer = app.query_one(MarkdownViewer)
        assert editor.has_focus
        assert editor.region.right <= app.query_one(AlignedPreview).region.x
        assert viewer.show_table_of_contents
        editor.load_text("# Updated\n")
        await pilot.pause()
        assert viewer.document.table_of_contents[0][1] == "Original"
        assert path.read_text() == "# Original\n"
        await pilot.press("escape")
        assert app.editing
        await pilot.press("ctrl+s", "escape")
        assert path.read_text() == "# Updated\n"
        assert not app.editing
        assert viewer.show_table_of_contents
        await pilot.press("e", "q", "r", "t", "e", "j", "k", "g")
        assert app.editing
        assert "qrtejkg" in editor.text
        await pilot.press("ctrl+d")
        assert not app.editing
        assert path.read_text() == "# Updated\n"
        assert viewer.document.table_of_contents[0][1] == "Updated"


async def test_save_failure_keeps_edits(tmp_path, monkeypatch):
    path = tmp_path / "edit.md"
    path.write_text("Original")
    app = Viewer(path)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.press("e")
        app.query_one(TextArea).load_text("Changed")
        monkeypatch.setattr(type(path), "write_text", Mock(side_effect=OSError("Read-only")))
        await pilot.press("ctrl+s", "escape")
        assert app.dirty
        assert app.query_one(TextArea).text == "Changed"
        assert path.read_text() == "Original"


async def test_edit_scroll_sync(tmp_path):
    path = tmp_path / "scroll.md"
    path.write_text("\n\n".join(f"## Section {i}\n\n" + "Some text. " * 20 for i in range(40)))
    app = Viewer(path, start_editing=True, live_edit=False)
    async with app.run_test(size=(100, 30)) as pilot:
        await app.workers.wait_for_complete()
        await settle_preview(app, pilot)
        editor = app.query_one(TextArea)
        preview = app.query_one(AlignedPreview)
        assert editor.max_scroll_y > 0
        assert preview.max_scroll_y > 0
        for source, target in ((editor, preview), (preview, editor)):
            for fraction in (1, 0):
                source.scroll_to(y=source.max_scroll_y * fraction, animate=False, immediate=True)
                await settle_preview(app, pilot)
                assert target.scroll_y == pytest.approx(target.max_scroll_y * fraction, abs=1)
                assert source.scroll_y == pytest.approx(source.max_scroll_y * fraction, abs=1)

        editor.move_cursor(editor.document.end)
        await settle_preview(app, pilot)
        assert preview.scroll_y == preview.max_scroll_y
        await pilot.press("enter", "x")
        await settle_preview(app, pilot)
        assert preview.scroll_y == pytest.approx(preview.max_scroll_y, abs=1)
        await pilot.resize_terminal(80, 24)
        await settle_preview(app, pilot)
        assert preview.scroll_y / preview.max_scroll_y == pytest.approx(
            editor.scroll_y / editor.max_scroll_y, abs=0.01
        )
        await pilot.press("ctrl+s", "escape")
        previous_editor_y = editor.scroll_y
        preview.scroll_to(y=0, animate=False, immediate=True)
        await settle_preview(app, pilot)
        assert editor.scroll_y == previous_editor_y


async def test_preview_rows_align_with_source(tmp_path):
    source = (
        "# First\n\n\n"
        "## Middle\n\n"
        + " ".join(f"word{i}" for i in range(100))
        + "\ncontinued **bold** text\n\n"
        "> Quote\n>\n> Another paragraph\n\n"
        "- First item\n- Second item\n\n"
        "```python\nprint('hello')\n```\n\n"
        "| A | B |\n| - | - |\n| 1 | 2 |\n\n"
        "## Last\n\n"
        "[repeat](https://example.com/repeat/repeat)repeat " + "你好 **world** " * 20
        + "\n\n```python\nvalue = '\t" + "long" * 45 + "'\n```\n\n"
        + "Tail paragraph\n\n" * 30
    )
    path = tmp_path / "blocks.md"
    path.write_text(source)
    app = Viewer(path, start_editing=True, live_edit=False)
    async with app.run_test(size=(100, 30)) as pilot:
        await app.workers.wait_for_complete()
        await settle_preview(app, pilot)
        editor = app.query_one(TextArea)
        preview = app.query_one(AlignedPreview)

        async def check_alignment():
            assert preview.content_region.y == editor.content_region.y
            assert len(preview.rows) == editor.wrapped_document.height
            for row_index, strip in enumerate(preview.rows):
                for segment in strip:
                    index = segment.style.meta.get(SOURCE) if segment.style else None
                    if index is not None:
                        location = preview.projection.location(index)
                        assert editor.wrapped_document.location_to_offset(location).y == row_index
            text = "".join(row.text for row in preview.rows)
            assert text.count("repeat") == 2
            assert "https://example.com" not in text
            assert "long" * 45 in text
            assert text.count("你好") == 20
            for line_index, line in enumerate(editor.document.lines):
                for column in range(len(line) + 1):
                    assert preview.projection.offset((line_index, column)).y == editor.wrapped_document.location_to_offset((line_index, column)).y
            for line_index, source_line in enumerate(editor.document.lines):
                for marker in ("First", "Middle", "continued", "Another", "Second", "print", "Last"):
                    if marker in source_line:
                        offset = preview.projection.offset((line_index, source_line.index(marker)))
                        assert marker in preview.rows[offset.y].text
            continuation = editor.document.lines.index("continued **bold** text")
            assert preview.projection.offset((continuation, 0)).y > preview.projection.offset((continuation - 1, 0)).y
            for source_pane, target in ((editor, preview), (preview, editor)):
                for fraction in (0, 0.37, 1):
                    source_pane.scroll_to(y=int(source_pane.max_scroll_y * fraction), animate=False, immediate=True)
                    await settle_preview(app, pilot)
                    assert target.scroll_y == source_pane.scroll_y

        await check_alignment()
        await pilot.resize_terminal(61, 24)
        await settle_preview(app, pilot)
        await check_alignment()
        editor.load_text("Intro\n\n" + editor.text)
        await settle_preview(app, pilot)
        await app.workers.wait_for_complete()
        await settle_preview(app, pilot)
        await check_alignment()
        await pilot.press("ctrl+s", "escape")
        assert not preview.display
        assert app.query_one(MarkdownViewer).display
        assert path.read_text() == "Intro\n\n" + source


async def test_edit_scroll_sync_short_document(tmp_path):
    path = tmp_path / "short.md"
    path.write_text("# Short\n")
    app = Viewer(path, start_editing=True)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        await pilot.press("pagedown", "pageup")
        assert app.query_one(TextArea).scroll_y == 0
        assert app.query_one(AlignedPreview).scroll_y == 0


async def test_aligned_preview_links_and_empty_edits(tmp_path, monkeypatch):
    path = tmp_path / "links.md"
    path.write_text("[Website](https://example.com)\n\n" + "Paragraph\n\n" * 30 + "## Target\n\n" + "Tail\n\n" * 30)
    browser = Mock(return_value=True)
    monkeypatch.setattr("mdv.app.webbrowser.open", browser)
    app = Viewer(path, start_editing=True, live_edit=False)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await settle_preview(app, pilot)
        editor = app.query_one(TextArea)
        preview = app.query_one(AlignedPreview)
        link = preview.projection.offset((0, 1))
        await pilot.click(preview, offset=(
            preview.content_region.x - preview.region.x + link.x,
            preview.content_region.y - preview.region.y + link.y,
        ))
        await settle_preview(app, pilot)
        browser.assert_called_once_with("https://example.com", new=2)
        preview.action_link("#target")
        await settle_preview(app, pilot)
        target_y = editor.wrapped_document.location_to_offset((62, 0)).y
        assert editor.scroll_y == target_y
        assert preview.scroll_y == preview.projection.offset((62, 0)).y
        editor.load_text("")
        await settle_preview(app, pilot)
        await app.workers.wait_for_complete()
        await settle_preview(app, pilot)
        assert all(not row.text.strip() for row in preview.rows)
        assert editor.scroll_y == preview.scroll_y == 0
        await pilot.press("ctrl+d")
        assert not app.editing
        await pilot.press("e")
        await settle_preview(app, pilot)
        assert "Website" in "".join(row.text for row in preview.rows)


async def test_converted_document_cannot_be_edited(tmp_path):
    path = tmp_path / "example.html"
    path.write_text("<h1>Hello</h1>")
    app = Viewer(path)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.press("e")
        assert not app.editing
        assert path.read_text() == "<h1>Hello</h1>"


def test_markdown_and_html(tmp_path):
    markdown = tmp_path / "example.MD"
    markdown.write_text("\ufeff# Hello\n", encoding="utf-8")
    assert load_document(markdown) == "# Hello\n"
    html = tmp_path / "example.html"
    html.write_text("<h1>Hello</h1><p>A <strong>bold</strong> word.</p>")
    converted = load_document(html)
    assert "# Hello" in converted
    assert "**bold**" in converted


def test_pptx(tmp_path):
    from pptx import Presentation

    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[1])
    slide.shapes.title.text = "Office document"
    slide.placeholders[1].text = "Conversion works."
    path = tmp_path / "example.pptx"
    document.save(path)
    assert "Conversion works." in load_document(path)


def test_cli(tmp_path, capsys, monkeypatch):
    path = tmp_path / "example.md"
    path.write_text("# Hello")
    assert main([str(path), "--raw"]) == 0
    assert capsys.readouterr().out == "# Hello\n"
    assert main([str(tmp_path / "missing.pdf")]) == 1
    assert "Not a file" in capsys.readouterr().err
    monkeypatch.setattr("sys.stdin", io.StringIO("# Piped\n"))
    assert main([]) == 0
    assert capsys.readouterr().out == "# Piped\n"


async def test_interaction_and_reload(tmp_path):
    path = tmp_path / "example.md"
    path.write_text("# First\n\n" + "Paragraph\n\n" * 100)
    app = Viewer(path)
    async with app.run_test(size=(100, 30)) as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        viewer = app.query_one(MarkdownViewer)
        assert viewer.document.table_of_contents[0][1] == "First"
        assert not viewer.show_table_of_contents
        await pilot.press("t")
        assert viewer.show_table_of_contents
        await pilot.press("G")
        await pilot.pause()
        assert viewer.scroll_y > 0
        await pilot.press("g")
        await pilot.pause()
        assert viewer.scroll_y == 0
        path.write_text("# Changed\n")
        await pilot.press("r")
        await app.workers.wait_for_complete()
        await pilot.pause()
        assert viewer.document.table_of_contents[0][1] == "Changed"
        path.unlink()
        await pilot.press("r")
        await app.workers.wait_for_complete()
        assert "Unable to load" in str(app.query_one("#status", Static).render())
        await pilot.press("ctrl+c")


async def test_link_click_does_not_trigger_default_file_navigation(tmp_path, monkeypatch):
    browser = Mock(return_value=True)
    monkeypatch.setattr("mdv.app.webbrowser.open", browser)
    path = tmp_path / "links.md"
    path.write_text("[Website](https://example.com)\n\n# Heading\n")
    app = Viewer(path)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        await pilot.pause()
        viewer = app.query_one(MarkdownViewer)
        navigate = AsyncMock()
        monkeypatch.setattr(viewer, "go", navigate)
        for href in ("https://example.com", "missing.md", "#heading"):
            viewer.document.post_message(Markdown.LinkClicked(viewer.document, href))
            await pilot.pause()
        navigate.assert_not_awaited()
        browser.assert_called_once_with("https://example.com", new=2)
        assert app.path == path


async def test_browser_failure_shows_error(tmp_path, monkeypatch):
    path = tmp_path / "links.md"
    path.write_text("# Links\n")
    browser = Mock(return_value=False)
    monkeypatch.setattr("mdv.app.webbrowser.open", browser)
    app = Viewer(path)
    async with app.run_test() as pilot:
        await app.workers.wait_for_complete()
        viewer = app.query_one(MarkdownViewer)
        notify = Mock()
        monkeypatch.setattr(viewer, "notify", notify)
        for failure in (None, OSError("Browser unavailable")):
            browser.side_effect = failure
            viewer.document.post_message(Markdown.LinkClicked(viewer.document, "https://example.com"))
            await pilot.pause()
            assert notify.call_args.kwargs["title"] == "Could not open browser"
            assert notify.call_args.kwargs["severity"] == "error"
